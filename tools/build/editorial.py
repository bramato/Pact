"""Build the illustrated Italian guide from its editable chapter source.

Authoring requires python-docx. --check-markdown uses only the standard library.
Render and inspect the DOCX with the documents skill before distributing it.
"""

import argparse
import html
import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "docs/editorial"
SOURCE = OUTPUT / "chapters.json"
DOCX = OUTPUT / "PACT-guida-illustrata.docx"
MARKDOWN = OUTPUT / "README.md"


def markdown(data):
    lines = [f"# {data['title']}", "", f"**Bozza {data['version']} · {data['date']}**", ""]
    lines += ["[Scarica il documento Word modificabile](PACT-guida-illustrata.docx)", ""]
    lines += [f'<img src="assets/00-overview.png" width="340" alt="{html.escape(data["sections"][0]["alt"], quote=True)}">', ""]
    for text in data["introduction"]:
        lines += [text, ""]
    lines += ["## Percorso di lettura", ""]
    for i, section in enumerate(data["sections"], 1):
        heading = f"{i} {section['title']}"
        anchor = re.sub(r"[^a-z0-9]+", "-", heading.lower()).strip("-")
        lines += [f"- [{heading}](#{anchor})"]
    lines += [""]
    for i, section in enumerate(data["sections"], 1):
        lines += [f"## {i} {section['title']}", "", section["lead"], ""]
        lines += [f'<img src="assets/{section["asset"]}" width="260" alt="{html.escape(section["alt"], quote=True)}">', "", f'*{section["caption"]}*', ""]
        for block in section["blocks"]:
            lines += [f"### {block['heading']}", ""]
            for text in block["paragraphs"]:
                lines += [text, ""]
            if block.get("code"):
                lines += ["```sh", *block["code"], "```", ""]
        lines += [f'[Approfondimento tecnico](../../{section["source"]})', ""]
    lines += ["## Glossario e riferimenti", "", "| Termine | Significato |", "| --- | --- |"]
    for term, meaning in data["glossary"]:
        lines += [f"| {term} | {meaning} |"]
    lines += ["", "I riferimenti seguenti appartengono al checkout della bozza. La specifica resta il riferimento per implementare il profilo; la guida ne spiega il percorso e l’uso.", ""]
    for ref in data["references"]:
        lines += [f'- [{ref["title"]}](../../{ref["path"]})']
    lines += ["", "Le illustrazioni sono disponibili come PNG con trasparenza reale. [Il catalogo dei prompt](art-direction.json) registra il sistema visivo e la revisione del grafo Swarm.", ""]
    return "\n".join(lines)


def document(data):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.85)
    section.footer_distance = Inches(0.35)
    section.different_first_page_header_footer = True
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
        borders = style.element.find(".//" + qn("w:pBdr"))
        if borders is not None:
            borders.getparent().remove(borders)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.widow_control = True
    for name, size, before, after in [("Title", 32, 0, 16), ("Subtitle", 15, 0, 18), ("Heading 1", 23, 0, 12), ("Heading 2", 14, 14, 6)]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    doc.styles["Caption"].font.size = Pt(9)
    doc.styles["Caption"].font.italic = True
    doc.styles["Caption"].paragraph_format.space_after = Pt(14)
    doc.core_properties.title = data["title"]
    doc.core_properties.subject = data["subtitle"]
    doc.core_properties.author = "Progetto PACT"
    doc.core_properties.keywords = "PACT A2A Core Swarm guida illustrata"
    doc.styles["Normal"].element.get_or_add_rPr().append(OxmlElement("w:lang"))
    doc.styles["Normal"].element.rPr[-1].set(qn("w:val"), "it-IT")

    def bookmark(paragraph, name, index):
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(index))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(index))
        paragraph._p.insert(0, start)
        paragraph._p.append(end)

    def link(paragraph, label, target=None, anchor=None):
        element = OxmlElement("w:hyperlink")
        if anchor:
            element.set(qn("w:anchor"), anchor)
        else:
            relation = doc.part.relate_to(target, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
            element.set(qn("r:id"), relation)
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "1268AE")
        properties.append(color)
        run.append(properties)
        text = OxmlElement("w:t")
        text.text = label
        run.append(text)
        element.append(run)
        paragraph._p.append(element)

    def picture(chapter, width):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        shape = p.add_run().add_picture(str(OUTPUT / "assets" / chapter["asset"]), width=Inches(width))
        shape._inline.docPr.set("descr", chapter["alt"])
        shape._inline.docPr.set("title", chapter["title"])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("PACT 0.2.0  ·  ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    doc.add_paragraph("PACT", "Subtitle")
    doc.add_paragraph(data["title"], "Title")
    doc.add_paragraph(data["subtitle"], "Subtitle")
    picture(data["sections"][0], 3.65)
    for text in data["introduction"]:
        doc.add_paragraph(text)
    doc.add_paragraph(f"Bozza {data['version']} · {data['date']}")
    doc.add_paragraph("Proposta indipendente sopra A2A · Licenza MIT")

    doc.add_page_break()
    doc.add_heading("Percorso di lettura", 1)
    doc.add_paragraph("La guida segue le undici sezioni della documentazione. Ogni tavola rappresenta una funzione; le spiegazioni collegano quella funzione al comportamento verificato della bozza.")
    for i, chapter in enumerate(data["sections"], 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(12)
        link(p, f"{i:02d}   {chapter['title']}", anchor="section_" + str(i))
    doc.add_heading("Un esempio comune", 2)
    doc.add_paragraph("Il testo PACT connects three careful agents attraversa il percorso di chiarimento, validazione e collaborazione. Il risultato contiene una sintesi, un conteggio di cinque parole e un’approvazione esplicita. La demo può essere ripetuta senza chiamare un modello AI.")

    for i, chapter in enumerate(data["sections"], 1):
        doc.add_page_break()
        heading = doc.add_heading(f"{i:02d} {chapter['title']}", 1)
        bookmark(heading, "section_" + str(i), i)
        doc.add_paragraph(chapter["lead"])
        picture(chapter, 2.05)
        p = doc.add_paragraph(chapter["caption"], "Caption")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for block in chapter["blocks"]:
            if block.get("page_break"):
                doc.add_page_break()
                doc.add_heading("Eventi e avanzamento nel Core", 1)
            doc.add_heading(block["heading"], 2)
            for text in block["paragraphs"]:
                doc.add_paragraph(text)
            for command in block.get("code", []):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(command)
                run.font.name = "Courier New"
                run.font.size = Pt(9)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        link(p, "Approfondimento tecnico nel checkout", "../../" + chapter["source"])

    doc.add_page_break()
    doc.add_heading("Glossario e riferimenti", 1)
    for term, meaning in data["glossary"]:
        p = doc.add_paragraph()
        p.add_run(term + "  ").bold = True
        p.add_run(meaning)
    doc.add_heading("Documenti di riferimento", 2)
    doc.add_paragraph("La specifica definisce il profilo; questa guida ne spiega l’uso. I riferimenti tecnici seguenti sono file del checkout della bozza e i collegamenti relativi funzionano quando il documento rimane nella cartella docs/editorial.")
    for ref in data["references"]:
        link(doc.add_paragraph(), ref["title"], "../../" + ref["path"])
    doc.save(DOCX)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check-markdown", action="store_true")
    args = parser.parse_args()
    data = json.loads(SOURCE.read_text())
    assert len(data["sections"]) == 11
    assert len({s["id"] for s in data["sections"]}) == 11
    for section in data["sections"]:
        assert (OUTPUT / "assets" / section["asset"]).is_file()
        assert (ROOT / section["source"]).is_file()
    text = markdown(data)
    if args.check_markdown:
        assert MARKDOWN.read_text() == text, "Regenerate editorial Markdown from chapters.json"
        print("Editorial source and Markdown agree; all 11 sections and assets are present")
        return
    MARKDOWN.write_text(text)
    document(data)
    print("Built editable DOCX and web guide for all 11 documentation sections")


if __name__ == "__main__":
    main()
